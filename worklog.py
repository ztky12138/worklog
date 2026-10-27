#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
worklog — 从 git 历史生成周报草稿

定位：一个仓库、多个仓库都行；噪声提交会被过滤；按目录自动归到模块名下。

零第三方依赖。想输出 docx 再加 python-docx（可选）。

用法：
    worklog week --repo . --since 2026-10-20 --until 2026-10-24
    worklog week --repo D:\\a --repo D:\\b --format md --out 周报.md
    worklog record --repo . --module 智慧法治

配置（worklog.json，可选）：
    {
      "repos": [{"path": ".", "name": "数字化平台"}],
      "modules": {"UI2/": "前端页面", "service/": "后端服务"},
      "ignore_commits": ["^Merge ", "^自动发布$", "^chore\\\\(release\\\\)"],
      "person": {"name": "", "role": ""}
    }
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import subprocess
import sys

CONFIG_NAME = 'worklog.json'

# 默认过滤：这些提交几乎不含"人做的事"，但会淹没周报
DEFAULT_IGNORE = [
    r'^Merge (branch|remote|pull request)',
    r'^Merge ',
    r'^自动发布',
    r'^chore\(release\)',
    r'^v?\d+\.\d+\.\d+$',          # 纯版本号
]

CREATE_NO_WINDOW = 0x08000000 if sys.platform == 'win32' else 0


# --------------------------------------------------------------------------
# 基础设施
# --------------------------------------------------------------------------
def setup_stdout() -> None:
    """Windows 控制台是 GBK，管道是 UTF-8 —— 按是否 tty 选编码。"""
    if sys.platform != 'win32':
        return
    try:
        enc = 'gbk' if sys.stdout.isatty() else 'utf-8'
        sys.stdout.reconfigure(encoding=enc, errors='replace')
        sys.stderr.reconfigure(encoding=enc, errors='replace')
    except Exception:
        pass


def run(cmd: list[str], cwd: str | None = None, timeout: int = 60) -> tuple[int, str, str]:
    try:
        p = subprocess.run(cmd, cwd=cwd, capture_output=True, timeout=timeout,
                           creationflags=CREATE_NO_WINDOW)
        return p.returncode, p.stdout.decode('utf-8', 'replace'), p.stderr.decode('utf-8', 'replace')
    except FileNotFoundError:
        return -1, '', 'git not found'
    except subprocess.TimeoutExpired:
        return -2, '', 'git timed out'
    except Exception as e:
        return -3, '', str(e)


def load_config(start: str) -> dict:
    """在起始目录及上级找 worklog.json，找不到就给空配置。"""
    cur = os.path.abspath(start)
    while True:
        p = os.path.join(cur, CONFIG_NAME)
        if os.path.isfile(p):
            try:
                with open(p, encoding='utf-8-sig') as f:
                    cfg = json.load(f)
                cfg['_path'] = p
                return cfg
            except Exception as e:
                print(f'[warn] 配置文件解析失败 {p}: {e}', file=sys.stderr)
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return {}


# --------------------------------------------------------------------------
# 时间
# --------------------------------------------------------------------------
def parse_date(s: str) -> dt.date:
    return dt.date.fromisoformat(s.strip())


def week_monday(d: dt.date) -> dt.date:
    return d - dt.timedelta(days=d.weekday())


def resolve_period(args) -> tuple[dt.date, dt.date]:
    today = dt.date.today()
    if getattr(args, 'last_week', False):
        mon = week_monday(today) - dt.timedelta(days=7)
        return mon, mon + dt.timedelta(days=6)
    since = parse_date(args.since) if args.since else week_monday(today)
    until = parse_date(args.until) if args.until else today
    if until < since:
        raise SystemExit(f'日期范围非法: {since} > {until}')
    return since, until


def period_label(since: dt.date, until: dt.date) -> str:
    if since.year == until.year:
        return f'{since:%Y-%m-%d} ~ {until:%m-%d}'
    return f'{since:%Y-%m-%d} ~ {until:%Y-%m-%d}'


# --------------------------------------------------------------------------
# git
# --------------------------------------------------------------------------
def is_git_repo(path: str) -> bool:
    return os.path.isdir(path) and os.path.isdir(os.path.join(path, '.git'))


def git_commits(repo: str, since: dt.date, until: dt.date) -> list[dict]:
    fmt = '%h\x1f%ad\x1f%an\x1f%s'
    rc, out, err = run(['git', '-C', repo, 'log',
                        f'--since={since} 00:00:00', f'--until={until} 23:59:59',
                        f'--pretty=format:{fmt}', '--date=short'])
    if rc != 0:
        print(f'[warn] git log 失败 {repo}: {err.strip()[:120]}', file=sys.stderr)
        return []
    commits = []
    for line in out.splitlines():
        parts = line.split('\x1f')
        if len(parts) < 4:
            continue
        commits.append({'hash': parts[0], 'date': parts[1],
                        'author': parts[2], 'subject': parts[3]})
    return commits


def commit_files(repo: str, sha: str) -> list[str]:
    rc, out, _ = run(['git', '-C', repo, 'show', '--name-only', '--pretty=format:', sha])
    if rc != 0:
        return []
    return [f.strip() for f in out.splitlines() if f.strip()]


# --------------------------------------------------------------------------
# 过滤 / 分组
# --------------------------------------------------------------------------
def compile_ignore(patterns: list[str]) -> list[re.Pattern]:
    out = []
    for p in patterns:
        try:
            out.append(re.compile(p, re.IGNORECASE))
        except re.error as e:
            print(f'[warn] 忽略规则无效 {p!r}: {e}', file=sys.stderr)
    return out


def is_noise(subject: str, patterns: list[re.Pattern]) -> bool:
    return any(p.search(subject) for p in patterns)


def module_of(files: list[str], modules: dict[str, str]) -> str:
    """把一次提交归到一个模块：按配置里的目录前缀匹配，没配上就退到顶层目录。

    配置的 key 是路径前缀（如 "UI2/"），value 是给人看的名字。
    """
    top = ''
    for f in files:
        norm = f.replace('\\', '/')
        if not top:
            parts = norm.split('/')
            top = parts[0] if len(parts) > 1 else '(根目录)'
        for prefix, name in modules.items():
            if norm.startswith(prefix.replace('\\', '/')):
                return name
    return modules.get(top + '/', top or '(未知)')


def group_commits(repo: str, commits: list[dict], modules: dict[str, str],
                  ignore: list[re.Pattern]) -> dict[str, list[dict]]:
    groups: dict[str, list[dict]] = {}
    for c in commits:
        if is_noise(c['subject'], ignore):
            continue
        name = module_of(commit_files(repo, c['hash']), modules)
        groups.setdefault(name, []).append(c)
    return groups


# --------------------------------------------------------------------------
# 输出
# --------------------------------------------------------------------------
def render_markdown(title: str, person: dict, since: dt.date, until: dt.date,
                    per_repo: list[tuple[str, dict[str, list[dict]]]],
                    total: int, skipped: int) -> str:
    lines = [f'# {title}（{period_label(since, until)}）', '']

    who = person.get('name') or ''
    role = person.get('role') or ''
    if who or role:
        lines.append(f'> {who}{"｜" if who and role else ""}{role}')
        lines.append('')

    lines.append('## 概览')
    lines.append('')
    lines.append(f'- 有效提交 **{total}** 笔（已过滤噪声 {skipped} 笔）')
    lines.append(f'- 覆盖 **{sum(len(g) for _, g in per_repo)}** 个模块分组')
    lines.append('')

    for repo_name, groups in per_repo:
        lines.append(f'## {repo_name}')
        lines.append('')
        if not groups:
            lines.append('_本周无有效提交_')
            lines.append('')
            continue
        for mod, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
            lines.append(f'### {mod}（{len(items)} 笔）')
            lines.append('')
            for c in sorted(items, key=lambda x: x['date']):
                lines.append(f'- `{c["hash"]}` {c["date"]} {c["subject"]}')
            lines.append('')

    lines.append('## 工时（待填）')
    lines.append('')
    lines.append('| 模块 | 工时 |')
    lines.append('|---|---|')
    for _, groups in per_repo:
        for mod, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
            lines.append(f'| {mod} | |')
    lines.append('')
    lines.append('## 下周计划（待填）')
    lines.append('')
    return '\n'.join(lines)


def render_docx(md_text: str, out_path: str) -> bool:
    """可选：把 markdown 转成简单 docx。没装 python-docx 就返回 False。"""
    try:
        from docx import Document
    except ImportError:
        return False

    doc = Document()
    for line in md_text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith('### '):
            doc.add_heading(s[4:], level=3)
        elif s.startswith('## '):
            doc.add_heading(s[3:], level=2)
        elif s.startswith('# '):
            doc.add_heading(s[2:], level=1)
        elif s.startswith('- '):
            doc.add_paragraph(s[2:], style='List Bullet')
        elif s.startswith('> '):
            p = doc.add_paragraph(s[2:])
            p.runs[0].italic = True
        elif s.startswith('|'):
            continue                      # 表格留到 Markdown 里看
        else:
            doc.add_paragraph(s)
    doc.save(out_path)
    return True


# --------------------------------------------------------------------------
# 命令
# --------------------------------------------------------------------------
def collect(args, cfg: dict) -> tuple[list[tuple[str, list[dict]]], list[re.Pattern], dict]:
    """按 --repo（可多次）或配置里的 repos 收集各仓库提交。"""
    modules = cfg.get('modules') or {}
    ignore = compile_ignore((cfg.get('ignore_commits') or []) + DEFAULT_IGNORE)

    repos = []
    if args.repo:
        for r in args.repo:
            repos.append({'path': r, 'name': os.path.basename(os.path.abspath(r))})
    elif cfg.get('repos'):
        repos = cfg['repos']
    else:
        repos = [{'path': '.', 'name': os.path.basename(os.getcwd())}]

    out = []
    for r in repos:
        path = r['path'] if os.path.isabs(r['path']) else os.path.join(
            os.path.dirname(cfg.get('_path', os.getcwd())), r['path'])
        path = os.path.abspath(path)
        if not is_git_repo(path):
            print(f'[warn] 不是 git 仓库，跳过: {path}', file=sys.stderr)
            continue
        out.append((r.get('name') or os.path.basename(path), path))
    return out, ignore, modules


def cmd_week(args, cfg: dict) -> int:
    since, until = resolve_period(args)
    repos, ignore, modules = collect(args, cfg)

    per_repo = []
    total = 0
    skipped = 0
    for name, path in repos:
        commits = git_commits(path, since, until)
        kept = [c for c in commits if not is_noise(c['subject'], ignore)]
        skipped += len(commits) - len(kept)
        groups = group_commits(path, commits, modules, ignore)
        total += len(kept)
        per_repo.append((name, groups))
        print(f'  {name}: {len(kept)} 笔有效 / {len(commits)} 笔总计')

    title = (cfg.get('title') or '工作周报')
    md = render_markdown(title, cfg.get('person') or {}, since, until,
                         per_repo, total, skipped)

    if args.out:
        out = args.out
    else:
        out = f'周报_{until:%Y%m%d}.md'

    if args.format == 'docx':
        if render_docx(md, out if out.endswith('.docx') else out.rsplit('.', 1)[0] + '.docx'):
            print(f'[OK] 已生成: {out.rsplit(".", 1)[0]}.docx')
        else:
            print('[warn] 没装 python-docx，退回 Markdown 输出', file=sys.stderr)
            with open(out if out.endswith('.md') else out.rsplit('.', 1)[0] + '.md',
                      'w', encoding='utf-8') as f:
                f.write(md)
            print(f'[OK] 已生成 Markdown')
    else:
        if not out.endswith('.md'):
            out = out.rsplit('.', 1)[0] + '.md'
        with open(out, 'w', encoding='utf-8') as f:
            f.write(md)
        print(f'[OK] 已生成: {out}')

    print(f'  共 {total} 笔有效提交，过滤 {skipped} 笔噪声')
    return 0


def cmd_record(args, cfg: dict) -> int:
    since, until = resolve_period(args)
    repos, ignore, modules = collect(args, cfg)
    if not repos:
        raise SystemExit('没有可用的 git 仓库')

    name, path = repos[0]
    commits = git_commits(path, since, until)
    kept = [c for c in commits if not is_noise(c['subject'], ignore)]
    groups = group_commits(path, kept, modules, ignore)

    module = args.module or name
    lines = [f'# {module} 开发记录', '',
             f'- 周期：{period_label(since, until)}',
             f'- 仓库：`{path}`',
             f'- 提交：{len(kept)} 笔（有效）', '']
    for mod, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        lines.append(f'## {mod}（{len(items)} 笔）')
        lines.append('')
        for c in sorted(items, key=lambda x: x['date']):
            lines.append(f'- `{c["hash"]}` {c["date"]} {c["subject"]}')
        lines.append('')

    out = args.out or f'{module}_开发记录_{until:%Y%m%d}.md'
    with open(out, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f'[OK] 已生成: {out}')
    return 0


# --------------------------------------------------------------------------
def main(argv=None) -> int:
    setup_stdout()
    ap = argparse.ArgumentParser(prog='worklog', description='从 git 历史生成周报草稿')
    sub = ap.add_subparsers(dest='cmd', required=True)

    def add_period(p):
        p.add_argument('--since', help='开始日期 YYYY-MM-DD（默认本周一）')
        p.add_argument('--until', help='结束日期 YYYY-MM-DD（默认今天）')
        p.add_argument('--last-week', action='store_true', help='改取上一周')

    p_w = sub.add_parser('week', help='生成周报草稿')
    add_period(p_w)
    p_w.add_argument('--repo', action='append', help='仓库路径，可多次')
    p_w.add_argument('--format', choices=['md', 'docx'], default='md')
    p_w.add_argument('--out', help='输出路径')

    p_r = sub.add_parser('record', help='生成开发记录')
    add_period(p_r)
    p_r.add_argument('--repo', action='append', help='仓库路径')
    p_r.add_argument('--module', help='模块名（用于标题和文件名）')
    p_r.add_argument('--out', help='输出路径')

    args = ap.parse_args(argv if argv is not None else sys.argv[1:])
    cfg = load_config(args.repo[0] if getattr(args, 'repo', None) else os.getcwd())

    try:
        if args.cmd == 'week':
            return cmd_week(args, cfg)
        if args.cmd == 'record':
            return cmd_record(args, cfg)
    except SystemExit:
        raise
    except Exception:
        import traceback
        print('[FAIL] ' + traceback.format_exc(), file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
